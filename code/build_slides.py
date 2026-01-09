#!/usr/bin/env python3
"""Build a PPTX deck for the term project."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BASE_DIR = Path(__file__).resolve().parents[1]
IMG_DIR = BASE_DIR / "images"
STOCK_DIR = IMG_DIR / "stock"
OUT_PATH = BASE_DIR / "Food_Price_Inflation_Turkey.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

COLORS = {
    "navy": RGBColor(31, 78, 121),
    "blue": RGBColor(28, 73, 120),
    "light_gray": RGBColor(245, 245, 245),
    "mid_gray": RGBColor(210, 210, 210),
    "dark_gray": RGBColor(60, 60, 60),
    "white": RGBColor(255, 255, 255),
}

STOCK_IMAGES = {
    "title": STOCK_DIR / "title_istanbul.jpg",
    "market": STOCK_DIR / "market_food.jpg",
    "currency": STOCK_DIR / "currency_exchange.jpg",
    "data": STOCK_DIR / "data_analysis.jpg",
}


def add_header(slide, title):
    bar_h = Inches(0.6)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, bar_h
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["light_gray"]
    bar.line.color.rgb = COLORS["light_gray"]

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, bar_h - Inches(0.05), SLIDE_W, Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["navy"]
    line.line.color.rgb = COLORS["navy"]

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(8), Inches(0.4))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLORS["dark_gray"]

    tag = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, SLIDE_W - Inches(1.8), Inches(0.12), Inches(1.4), Inches(0.3)
    )
    tag.fill.solid()
    tag.fill.fore_color.rgb = COLORS["navy"]
    tag.line.color.rgb = COLORS["navy"]
    tag_tf = tag.text_frame
    tag_tf.text = "DSA 210"
    tag_tf.paragraphs[0].font.size = Pt(10)
    tag_tf.paragraphs[0].font.color.rgb = COLORS["white"]
    tag_tf.paragraphs[0].alignment = PP_ALIGN.CENTER


def add_bullets(slide, left, top, width, height, bullets, font_size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, item in enumerate(bullets):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = COLORS["dark_gray"]


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_img = STOCK_IMAGES["title"]
    if title_img.exists():
        slide.shapes.add_picture(str(title_img), 0, 0, width=SLIDE_W, height=SLIDE_H)
        overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        overlay.fill.solid()
        overlay.fill.fore_color.rgb = COLORS["navy"]
        overlay.fill.transparency = 0.25
        overlay.line.color.rgb = COLORS["navy"]
    else:
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLORS["navy"]
        bg.line.color.rgb = COLORS["navy"]

    title = slide.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.2))
    tf = title.text_frame
    tf.text = "DSA 210 TERM PROJECT"
    p = tf.paragraphs[0]
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]

    subtitle = slide.shapes.add_textbox(Inches(0.9), Inches(2.9), Inches(11.5), Inches(1.0))
    tf = subtitle.text_frame
    tf.text = "Food Price Inflation in Turkey (2019-2024)"
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.font.color.rgb = COLORS["white"]

    meta = slide.shapes.add_textbox(Inches(0.9), Inches(4.1), Inches(11.5), Inches(0.8))
    tf = meta.text_frame
    tf.text = "Isik Giray Onal | 34088"
    p = tf.paragraphs[0]
    p.font.size = Pt(16)
    p.font.color.rgb = COLORS["white"]


def add_section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "")
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.8), Inches(11), Inches(1))
    tf = box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLORS["navy"]
    return slide


def add_text_slide(prs, title, bullets, image_path=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, title)
    if image_path and image_path.exists():
        img_w = Inches(4.2)
        img_h = Inches(3.0)
        img_x = SLIDE_W - img_w - Inches(0.6)
        img_y = Inches(2.0)
        slide.shapes.add_picture(str(image_path), img_x, img_y, width=img_w, height=img_h)
        text_w = Inches(7.6)
    else:
        text_w = Inches(12)
    add_bullets(slide, Inches(0.8), Inches(1.4), text_w, Inches(5.5), bullets)
    return slide


def add_image_text_slide(prs, title, image_path, bullets, img_left=True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, title)

    img_w = Inches(7.3)
    img_h = Inches(4.6)
    img_x = Inches(0.7) if img_left else Inches(5.3)
    img_y = Inches(1.4)
    slide.shapes.add_picture(str(image_path), img_x, img_y, width=img_w, height=img_h)

    text_x = Inches(8.2) if img_left else Inches(0.7)
    add_bullets(slide, text_x, Inches(1.4), Inches(4.5), Inches(4.6), bullets, font_size=16)
    return slide


def build_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_title_slide(prs)

    add_text_slide(
        prs,
        "Introduction",
        [
            "Analyze monthly food price indices (bread, milk, meat) and USD/TRY from 2019-2024.",
            "Construct a composite Food Price Index to summarize food inflation dynamics.",
            "Use EDA and correlation tests to quantify the currency-food link.",
            "Evaluate basic time-series regression models for short-term prediction.",
        ],
        image_path=STOCK_IMAGES["market"],
    )

    add_text_slide(
        prs,
        "Motivation",
        [
            "Food inflation directly affects household purchasing power.",
            "USD/TRY depreciation raises import and input costs for food producers.",
            "Understanding co-movement helps explain macro-level inflation pressures.",
        ],
        image_path=STOCK_IMAGES["currency"],
    )

    add_text_slide(
        prs,
        "Research Questions",
        [
            "How strongly do food price indices co-move with USD/TRY?",
            "Which food group (bread, milk, meat) shows the fastest growth?",
            "Do lag and rolling features improve short-term prediction?",
        ],
        image_path=STOCK_IMAGES["data"],
    )

    add_text_slide(
        prs,
        "Data Collection and Preparation",
        [
            "USD/TRY: TCMB EVDS monthly series (TP.DK.USD.A.YTL).",
            "Food indices: Eurostat HICP via FRED (2015=100).",
            "Align monthly dates, compute composite Food Price Index.",
            "Engineer lag and rolling features for modeling.",
        ],
        image_path=STOCK_IMAGES["market"],
    )

    add_section_slide(prs, "Exploratory Data Analysis")

    add_image_text_slide(
        prs,
        "Analysis",
        IMG_DIR / "Figure_1.png",
        [
            "Both series trend upward over 2019-2024.",
            "Strong co-movement suggests currency pass-through.",
            "Composite index smooths item-level noise.",
        ],
    )

    add_image_text_slide(
        prs,
        "Analysis (cont.)",
        IMG_DIR / "Figure_2.png",
        [
            "Pearson r = 0.9891 (p = 5.26e-60).",
            "Relationship is strongly positive and significant.",
            "Supports H1: currency and food inflation co-move.",
        ],
    )

    add_image_text_slide(
        prs,
        "Analysis (cont.)",
        IMG_DIR / "Figure_3.png",
        [
            "Meat index grows fastest (~10.3x).",
            "Bread (~8.0x) and milk (~7.7x) rise steadily.",
            "Composite tracks the shared upward trend.",
        ],
    )

    add_image_text_slide(
        prs,
        "Analysis (cont.)",
        IMG_DIR / "Figure_4.png",
        [
            "MoM changes show clustered volatility.",
            "Currency shocks align with larger food jumps.",
            "Volatility increases in later years.",
        ],
    )

    add_image_text_slide(
        prs,
        "Machine Learning Applications",
        IMG_DIR / "prediction_plot.png",
        [
            "Time-aware CV with lag and rolling features.",
            "Best model: Ridge (RMSE 45.15, R2 -0.025).",
            "Linear models underfit; richer features needed.",
        ],
    )

    add_text_slide(
        prs,
        "Conclusion",
        [
            "Food prices and USD/TRY move closely together.",
            "Meat prices show the fastest index growth.",
            "Short-term regression remains challenging without richer signals.",
        ],
        image_path=STOCK_IMAGES["title"],
    )

    add_text_slide(
        prs,
        "Sources",
        [
            "TCMB EVDS: https://evds2.tcmb.gov.tr",
            "FRED (Eurostat HICP): https://fred.stlouisfed.org",
            "Series: CP0111TRM086NEST, CP0112TRM086NEST, CP0114TRM086NEST",
            "Stock images: Wikimedia Commons (Istanbul skyline 02, Food stall at Christmas market, MICEX, Wikidata statistics)",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS["navy"]
    bg.line.color.rgb = COLORS["navy"]
    box = slide.shapes.add_textbox(Inches(0), Inches(3.0), SLIDE_W, Inches(1))
    tf = box.text_frame
    tf.text = "Thank you"
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(32)
    p.font.color.rgb = COLORS["white"]

    prs.save(OUT_PATH)
    print(f"Saved {OUT_PATH}")


if __name__ == "__main__":
    build_deck()
